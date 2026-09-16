"""Step 5a - Build the workbook shell with openpyxl (all sheets, formulas, charts).
Step 5b (06_workbook_com.ps1) then adds the Power Query queries, loads Source_Data and
Source_Durations, recalculates in Excel and scans for errors.

Formulas reference Source_Data / Source_Durations by full column so they resolve once
Power Query has loaded the tables.
"""
import json
import pathlib
import shutil
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo
from config import GROUPS, SHORT, MAPPING, ADMIN_RULE_TEXT, CANDIDATES, FOCUS_FAMILY, G1, G2, G3, G5, AUTHOR

ROOT = pathlib.Path(__file__).resolve().parents[1]
R = json.loads((ROOT / "build" / "results.json").read_text())
META = json.loads((ROOT / "data" / "raw" / "dataset_metadata.json").read_text())
OUTD = ROOT / "outputs"
DATA_OUT = OUTD / "data"
DATA_OUT.mkdir(parents=True, exist_ok=True)
STAGE = ROOT / "build" / "workbook_stage1.xlsx"
N = R["profile"]["rows"]
EXTRACT = R["extraction"]["extracted_at_utc"]

# ------------------------------------------------------------------ data files for Power Query
clean_src = ROOT / "data" / "processed" / "cohort_2025_clean.csv"
shutil.copyfile(clean_src, DATA_OUT / "cohort_2025_clean.csv")
shutil.copyfile(ROOT / "data" / "processed" / "closure_outcome_mapping.csv", DATA_OUT / "closure_outcome_mapping.csv")
c = pd.read_csv(clean_src, dtype={"days_to_close": "Int64"}, keep_default_na=False, na_values=[""])
c["outcome_group"] = c["closure_reason"].map(lambda r: MAPPING[r][0])
e = c[c["duration_eligible"] == "Yes"].copy()
fam_dept, fam_types = CANDIDATES[FOCUS_FAMILY]
views = []
for pop, sub in (("All published records", e), ("Excluding admin/internal-sounding", e[e["admin_sounding"] == "No"])):
    views.append(sub.assign(population=pop, scope="Citywide", request_type="All types", outcome_key=sub["outcome_group"]))
    views.append(sub.assign(population=pop, scope="Citywide", request_type="All types", outcome_key="All outcomes (mixed)"))
    f = sub[sub["focus_family"] == "Yes"]
    for tname, ft in [("All three types", f)] + [(t, f[f["service_request_type"] == t]) for t in fam_types]:
        views.append(ft.assign(population=pop, scope="Focus family", request_type=tname, outcome_key=ft["outcome_group"]))
        views.append(ft.assign(population=pop, scope="Focus family", request_type=tname, outcome_key="All outcomes (mixed)"))
        for r in ("Further action has been planned", "Referred to another service group"):
            fr = ft[ft["closure_reason"] == r]
            views.append(fr.assign(population=pop, scope="Focus family", request_type=tname, outcome_key=r))
V = pd.concat(views, ignore_index=True)
keys = ["population", "scope", "request_type", "outcome_key"]
freq = V.groupby(keys + ["days_to_close"]).size().rename("records").reset_index().sort_values(keys + ["days_to_close"])
freq["cum_records"] = freq.groupby(keys)["records"].cumsum()
freq.to_csv(DATA_OUT / "duration_frequency_2025.csv", index=False)
src_rows = c.groupby(["open_month", "department", "service_request_type", "channel", "status", "closure_reason",
                      "outcome_group", "admin_sounding", "focus_family", "local_area", "duration_exclusion_reason",
                      "duplicate_looking"], dropna=False).ngroups
print("expected Source_Data rows:", src_rows, "| duration frequency rows:", len(freq))

# ------------------------------------------------------------------ styles
FONT = "Arial"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=16, bold=True, color="17324D")
f_h2 = Font(name=FONT, size=12, bold=True, color="17324D")
f_note = Font(name=FONT, size=9, italic=True, color="52514E")
f_input = Font(name=FONT, size=10, color="0000FF")
f_hdr = Font(name=FONT, size=10, bold=True, color="FFFFFF")
fill_hdr = PatternFill("solid", fgColor="17324D")
fill_sec = PatternFill("solid", fgColor="E8EEF4")
fill_check = PatternFill("solid", fgColor="EAF5EA")
thin = Side(style="thin", color="C3C2B7")
b_all = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top")
PCT, INT = "0.0%", "#,##0"

wb = Workbook()
ws_readme = wb.active
ws_readme.title = "README"
order = ["Source_Data", "Source_Durations", "Data_Dictionary", "Closure_Outcome_Mapping", "Data_Quality_Log",
         "Analysis", "Dashboard", "Assumptions_and_Limitations"]
S = {"README": ws_readme}
for name in order:
    S[name] = wb.create_sheet(name)
for ws in S.values():
    ws.sheet_view.showGridLines = False

def put(ws, ref, value, font=f_base, fmt=None, fill=None, align=None, border=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    return cell

def header_row(ws, row, col, labels, widths=None):
    for i, lab in enumerate(labels):
        put(ws, f"{get_column_letter(col + i)}{row}", lab, f_hdr, fill=fill_hdr, align=Alignment(wrap_text=True, vertical="center"), border=b_all)
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(col + i)].width = w

def add_table(ws, ref, name, style="TableStyleMedium2"):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)

# Source_Data column letters (Power Query output order, table header in row 3)
SD_COLS = ["open_month", "department", "service_request_type", "channel", "status", "closure_reason",
           "outcome_group", "admin_sounding", "focus_family", "local_area", "duration_exclusion_reason",
           "duplicate_looking", "records"]
SDL = {name: get_column_letter(i + 1) for i, name in enumerate(SD_COLS)}
def col(name):
    L = SDL[name]
    return f"Source_Data!${L}:${L}"
def SUMIFS(*crit):
    if not crit:
        return f"=SUM({col('records')})"
    parts = ",".join(f"{col(k)},{v}" for k, v in crit)
    return f"=SUMIFS({col('records')},{parts})"
def q(s):
    return '"' + s.replace('"', '""') + '"'

DU_COLS = ["population", "scope", "request_type", "outcome_key", "days_to_close", "records", "cum_records"]
DUL = {name: get_column_letter(i + 1) for i, name in enumerate(DU_COLS)}
def dcol(name):
    L = DUL[name]
    return f"Source_Durations!${L}:${L}"

# ------------------------------------------------------------------ Source sheets (placeholders)
for name, desc in (("Source_Data", "Power Query output 'Source_Data': the cleaned 2025 cohort grouped to one row per unique combination of fields, with a record count. Outcome groups come from the Closure_Outcome_Mapping table during refresh."),
                   ("Source_Durations", "Power Query output 'Source_Durations': frequency of recorded calendar days to closure for duration-eligible closed cases, with cumulative counts used for nearest-rank percentiles.")):
    put(S[name], "A1", desc, f_note)

# ------------------------------------------------------------------ README
ws = ws_readme
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110
put(ws, "A1", "What Does Closed Mean? Making Vancouver 3-1-1 Performance Reporting Decision-Ready", f_title)
put(ws, "A2", "Independent, outside-in case study using public data. Not affiliated with, endorsed by, or informed by internal knowledge of the City of Vancouver.", f_note)
rows = [
    ("Purpose", "Supporting workbook for the four-page case study. It shows how the public 3-1-1 data was profiled, how closure reasons were mapped to analyst-defined outcome groups, and how every figure in the report was calculated."),
    ("Main question", "How can Vancouver's public 3-1-1 data be interpreted and reported so that 'closed' cases are not automatically mistaken for completed services?"),
    ("Data source", "City of Vancouver Open Data Portal, dataset '3-1-1 service requests': https://opendata.vancouver.ca/explore/dataset/3-1-1-service-requests/ (Open Government Licence - Vancouver)."),
    ("Extraction", f"API export downloaded {EXTRACT} (UTC). Dataset last modified {R['extraction']['dataset_modified']} per portal metadata."),
    ("API filter used", R["extraction"]["api_where_clause"] + "  (timezone=UTC). A buffered window; the 2025 cohort is then filtered locally."),
    ("Cohort definition", f"Requests whose opening timestamp, converted from UTC to America/Vancouver, falls on a 2025 local date: {N:,} records."),
    ("Privacy", "Address, latitude, longitude and geometry fields were removed before any file in this package was created. Only yes/no presence flags and local-area names are kept."),
    ("How to refresh", "All figures are already saved in this file, so no refresh is needed to read it. To rebuild from the CSV files: 1) Keep the 'data' folder next to this workbook (or type new full paths in the blue cells below). 2) Data > Refresh All. Power Query re-reads the CSV files, re-applies the Closure_Outcome_Mapping table and reloads Source_Data and Source_Durations. All Analysis and Dashboard figures are formulas and update automatically."),
    ("Colour conventions", "Blue text = typed inputs or lists chosen by the analyst. Black = formulas. Green-shaded rows on Analysis = reconciliation checks (should read TRUE)."),
    ("Percentile method", "Nearest-rank: the p-th percentile is the smallest number of recorded calendar days within which at least p% of eligible closed cases were closed."),
    ("Author", f"{AUTHOR} - portfolio case study prepared for a Business Analyst application."),
]
r = 4
for k, v in rows:
    put(ws, f"A{r}", k, f_bold, align=wrap)
    put(ws, f"B{r}", v, f_base, align=wrap)
    r += 1
r += 1
put(ws, f"A{r}", "Power Query file paths", f_h2); r += 1
put(ws, f"A{r}", "Workbook folder (formula)", f_bold)
put(ws, f"B{r}", '=LEFT(CELL("filename",A1),FIND("[",CELL("filename",A1))-1)', f_base); folder_row = r; r += 1
put(ws, f"A{r}", "Cleaned cohort CSV", f_bold)
put(ws, f"B{r}", f'=B{folder_row}&"data\\cohort_2025_clean.csv"', f_input); data_row = r; r += 1
put(ws, f"A{r}", "Duration frequency CSV", f_bold)
put(ws, f"B{r}", f'=B{folder_row}&"data\\duration_frequency_2025.csv"', f_input); dur_row = r; r += 1
wb.defined_names["DataFile"] = DefinedName("DataFile", attr_text=f"README!$B${data_row}")
wb.defined_names["DurationFile"] = DefinedName("DurationFile", attr_text=f"README!$B${dur_row}")
r += 1
put(ws, f"A{r}", "Sheet guide", f_h2); r += 1
guide = [
    ("Source_Data", "Power Query output: grouped cohort (one row per unique field combination + record count)."),
    ("Source_Durations", "Power Query output: days-to-closure frequency table for eligible closed cases."),
    ("Data_Dictionary", "Every published and derived field, format, missingness and open questions."),
    ("Closure_Outcome_Mapping", "All 17 original closure-reason values mapped to six analyst-defined outcome groups, with rationale and validation questions. Drives Power Query."),
    ("Data_Quality_Log", "Issue found > rule applied > records affected > effect on analysis."),
    ("Analysis", "All calculations as formulas, with reconciliation checks."),
    ("Dashboard", "Management view: headline measures and four charts."),
    ("Assumptions_and_Limitations", "What the public data can and cannot support, and what to validate with City staff."),
]
for k, v in guide:
    put(ws, f"A{r}", k, f_bold); put(ws, f"B{r}", v, f_base, align=wrap); r += 1

# ------------------------------------------------------------------ Data_Dictionary
ws = S["Data_Dictionary"]
put(ws, "A1", "Data dictionary - published fields and analyst-derived fields", f_title)
put(ws, "A2", "Published descriptions are from the City's dataset metadata. Missing counts are for the 2025 cohort (" + f"{N:,}" + " records).", f_note)
hdr = ["Field", "Source", "Format in extract", "Description", "How used in this study", "Missing (2025)", "Missing %", "Question to resolve / note"]
header_row(ws, 4, 1, hdr, [30, 13, 26, 46, 44, 12, 10, 50])
meta_desc = {f["name"]: (f.get("description") or "").strip() for f in META.get("fields", [])}
miss = {m["field"]: m["missing"] for m in R["profile"]["missing"]}
dd = [
    ("department", "Published", "Text", "How: compare demand and outcome mixes by organisational unit.", "Department names change over time; 'ZZ - OLD' prefixes mark obsolete units."),
    ("service_request_type", "Published", "Text", "Demand ranking; focus-family definition; admin/internal-sounding flag.", "Some type names sound administrative or internal; confirm which are resident-initiated."),
    ("status", "Published", "Text: 'Close' or 'Open'", "Separates closed cases (outcome analysis) from open cases (reported separately).", "Last known status only; no history of reopening is published."),
    ("closure_reason", "Published", "Text (17 values in 2025)", "Mapped to analyst-defined outcome groups; original value always preserved.", "Operational meaning of each value must be confirmed; 'N/A' appears on open cases."),
    ("service_request_open_timestamp", "Published", "ISO 8601 with +00:00 (UTC)", "Converted to America/Vancouver to derive the local opening date and month.", "Confirmed as true UTC: local-time conversion puts 70% of requests between 9:00 and 16:59."),
    ("service_request_close_date", "Published", "Date only (YYYY-MM-DD)", "Recorded calendar days to closure = close date minus local opening date.", "No time of day, so durations are calendar days, not elapsed time. Evidence indicates a local-calendar date."),
    ("last_modified_timestamp", "Published", "ISO 8601 with +00:00 (UTC)", "Consistency checks only. Not used as a closure time.", "May change after closure for reasons unrelated to service."),
    ("address", "Published", "Text", "Presence flag only (privacy). Values removed.", "Withheld for some request types by design (City metadata)."),
    ("local_area", "Published", "Text (22 local planning areas)", "Aggregated geographic summaries only.", "Blank for many non-location request types such as feedback cases."),
    ("channel", "Published", "Text (10 values)", "Intake channel comparisons within comparable request types.", "Definitions of 'Unknown', 'Mail Out' and 'Face To Face' channels are not published."),
    ("latitude", "Published", "Decimal", "Presence flag only (privacy). Values removed.", "Some addresses do not geocode (City metadata)."),
    ("longitude", "Published", "Decimal", "Presence flag only (privacy). Values removed.", ""),
    ("geom", "Published", "Geo point", "Not used.", ""),
]
r = 5
for name, src, fmt, how, note in dd:
    vals = [name, src, fmt, meta_desc.get(name, ""), how, miss.get(name, 0), miss.get(name, 0) / N, note]
    for i, v in enumerate(vals):
        cell = put(ws, f"{get_column_letter(i + 1)}{r}", v, f_base, align=wrap, border=b_all)
    ws[f"F{r}"].number_format = INT
    ws[f"G{r}"].number_format = PCT
    r += 1
derived = [
    ("open_datetime_local / open_date_local", "Derived", "Local datetime / date", "Opening timestamp converted from UTC to America/Vancouver (daylight saving handled).", "Defines the 2025 cohort and monthly demand."),
    ("open_month", "Derived", "Integer 1-12", "Local opening month.", "Monthly demand and outcome trends."),
    ("days_to_close", "Derived", "Integer (days)", "close_date minus open_date_local. Blank for open cases.", "Recorded calendar days to closure. Not service-delivery or resolution time."),
    ("duration_eligible / duration_exclusion_reason", "Derived", "Yes/No; text", "Eligible = closed, has close date, and days_to_close >= 0.", "Excluded rows are counted on Data_Quality_Log and Analysis."),
    ("outcome_group / outcome_subgroup", "Derived (analyst-defined)", "Text", "Interpretation of closure_reason via Closure_Outcome_Mapping.", "Requires validation with service owners."),
    ("admin_sounding", "Derived", "Yes/No", ADMIN_RULE_TEXT, "Name-based only. Does not prove a record is internal."),
    ("zz_old_label", "Derived", "Yes/No", "Department or request type begins with 'ZZ'.", "City metadata: 'ZZ - OLD' marks obsolete types."),
    ("duplicate_looking", "Derived", "Yes/No", "All published fields identical to another record except last-modified timestamp.", "Cannot be confirmed as duplicates without a case identifier."),
    ("focus_family", "Derived", "Yes/No", f"{FOCUS_FAMILY}: {fam_dept} - " + ", ".join(fam_types) + ".", "Selected on evidence (see Analysis)."),
    ("has_address / has_coordinates", "Derived", "Yes/No", "Presence of the withheld location fields.", "Supports missingness analysis without publishing locations."),
    ("extract_row", "Derived", "Integer", "Row number in the saved extract.", "NOT a City case identifier. None is published."),
]
for name, src, fmt, desc, note in derived:
    vals = [name, src, fmt, desc, note, None, None, ""]
    for i, v in enumerate(vals):
        put(ws, f"{get_column_letter(i + 1)}{r}", v, f_base, align=wrap, border=b_all)
    r += 1
ws.freeze_panes = "B5"

# ------------------------------------------------------------------ Closure_Outcome_Mapping
ws = S["Closure_Outcome_Mapping"]
put(ws, "A1", "Closure-reason mapping (analyst-defined interpretation)", f_title)
put(ws, "A2", "Every original closure-reason value is preserved and mapped to one outcome group. This is an outside-in interpretation of published wording, not a City definition. Power Query uses this table (tblClosureMapping) during refresh.", f_note)
mh = ["Closure reason (original)", "Outcome group (analyst-defined)", "Outcome sub-group", "Closed cases 2025", "Open cases 2025", "Share of closed cases", "Rationale for grouping", "Question to validate with City staff"]
header_row(ws, 4, 1, mh, [40, 42, 30, 14, 12, 12, 60, 64])
r = 5
for reason, (g, sg, why, qn) in MAPPING.items():
    put(ws, f"A{r}", reason, f_input, border=b_all)
    put(ws, f"B{r}", g, f_input, border=b_all)
    put(ws, f"C{r}", sg, f_input, border=b_all)
    put(ws, f"D{r}", SUMIFS(("status", q("Close")), ("closure_reason", f"$A{r}")), f_base, INT, border=b_all)
    put(ws, f"E{r}", SUMIFS(("status", q("Open")), ("closure_reason", f"$A{r}")), f_base, INT, border=b_all)
    put(ws, f"F{r}", f"=IFERROR(D{r}/SUM($D$5:$D${4 + len(MAPPING)}),0)", f_base, PCT, border=b_all)
    put(ws, f"G{r}", why, f_base, align=wrap, border=b_all)
    put(ws, f"H{r}", qn, f_base, align=wrap, border=b_all)
    r += 1
last_map = r - 1
add_table(ws, f"A4:H{last_map}", "tblClosureMapping")
r += 1
put(ws, f"A{r}", "Outcome-group summary (closed cases)", f_h2); r += 1
header_row(ws, r, 1, ["Outcome group", "Closed cases", "Share of closed", "Values mapped"]); r += 1
grp_start = r
for g in GROUPS:
    put(ws, f"A{r}", g, f_base, border=b_all)
    put(ws, f"B{r}", f'=SUMIFS($D$5:$D${last_map},$B$5:$B${last_map},A{r})', f_base, INT, border=b_all)
    put(ws, f"C{r}", f"=IFERROR(B{r}/SUM($B${grp_start}:$B${grp_start + 5}),0)", f_base, PCT, border=b_all)
    put(ws, f"D{r}", f'=COUNTIF($B$5:$B${last_map},A{r})', f_base, border=b_all)
    r += 1
put(ws, f"A{r}", "Check: all closure reasons found in Source_Data are mapped", f_bold, fill=fill_check)
put(ws, f"B{r}", f'=SUMIFS({col("records")},{col("outcome_group")},"*UNMAPPED*")=0', f_bold, fill=fill_check)
map_check = f"Closure_Outcome_Mapping!B{r}"
ws.freeze_panes = "B5"

# ------------------------------------------------------------------ Data_Quality_Log
ws = S["Data_Quality_Log"]
put(ws, "A1", "Data quality log - issue found > rule applied > records affected > effect on analysis", f_title)
put(ws, "A2", "Counts shown as formulas are recalculated from Source_Data. Counts marked 'Python' were computed from the row-level extract by 03_analyze.py because the underlying fields are not loaded into this workbook (for example, withheld location fields).", f_note)
qh = ["ID", "Check", "Field(s)", "Records affected", "% of cohort", "Finding", "Rule applied", "Effect on analysis", "Validation needed", "Computed by"]
header_row(ws, 4, 1, qh, [7, 30, 24, 13, 10, 52, 44, 44, 44, 11])
Q = R["quality"]
DB = R["date_behaviour"]
dq = [
    ("Row count and cohort", "open timestamp", SUMIFS(), f"{N:,} records opened on a 2025 Vancouver-local date. The UTC-year count is {R['cohort_reconciliation']['utc_year_2025_rows']:,}; the local-date cohort adds {R['cohort_reconciliation']['local_2025_but_utc_2026']} records opened on the evening of 31 Dec 2025 (local) and removes {R['cohort_reconciliation']['utc_2025_but_local_2024']} opened on the evening of 31 Dec 2024 (local).", "Filter on local opening date after UTC to America/Vancouver conversion.", "All analysis uses the local-date cohort.", "None.", "Formula"),
    ("Time zone of opening timestamp", "service_request_open_timestamp", None, f"All values carry +00:00. Converted to local time, {DB['peak_local_hours_share_9_to_16']:.0%} of requests open between 9:00 and 16:59 and {DB['share_opened_0_to_6_local']:.1%} between midnight and 6:59, consistent with true UTC.", "Convert to America/Vancouver before deriving dates.", "Local dates are used consistently.", "Confirm the source system's time-zone handling.", "Python"),
    ("Close-date calendar", "service_request_close_date", None, f"Date only, no time. For {DB['evening_auto_closed_n']:,} evening (17:00+) auto-closed cases, the close date equals the local opening date {DB['evening_auto_closed_close_eq_local_open_date']:.1%} of the time and the UTC date {DB['evening_auto_closed_close_eq_utc_open_date']:.1%}: evidence of a local-calendar date.", "Treat as a local calendar date. Report 'recorded calendar days to closure'.", "Durations are whole calendar days, not elapsed time.", "Confirm how the close date is stamped.", "Python"),
    ("Close date before local opening date", "close date, open date", SUMIFS(("duration_exclusion_reason", q("Close date before local open date"))), f"{Q['negative_days']} closed cases have a close date one day before the local opening date. All are 'Referred to another service group' and opened between 6:00 and 7:59 local time.", "Not replaced with zero. Excluded from duration statistics; kept in counts.", "Negligible effect on percentiles.", "Ask how referral close dates are stamped.", "Formula"),
    ("Open cases (no close date)", "status, close date", SUMIFS(("status", q("Open"))), f"Open at extraction, all opened in 2025. Aged {R['open_cases']['age_min']}-{R['open_cases']['age_max']} days at extraction (median {R['open_cases']['age_p50']}).", "Reported separately. Excluded from closed-case durations.", "Closed-case durations do not describe these cases.", "Are these genuinely active, or awaiting administrative closure?", "Formula"),
    ("Closed without a close date", "status, close date", None, f"{Q['closed_without_close_date']} records.", "None required.", "None.", "None.", "Python"),
    ("Status and closure-reason combinations", "status, closure_reason", SUMIFS(("status", q("Open")), ("closure_reason", q("<>N/A"))), f"'N/A' appears only on open cases ({Q['open_na_reason']:,}). {Q['open_with_non_na_reason']} open cases carry a closure reason (for example 'Referred to another service group').", "Outcome shares use closed cases only. Open cases with a reason are flagged, not re-coded.", "Small; flagged as an exception type for the proposed report.", "Can a case be reopened while keeping its earlier reason?", "Formula"),
    ("Closed cases recorded as 'N/A'", "status, closure_reason", SUMIFS(("status", q("Close")), ("closure_reason", q("N/A"))), "None: 'N/A' behaves as the placeholder for open cases in 2025.", "Grouped with 'Unknown' as an unclear outcome.", "None for closed-case shares.", "Confirm 'N/A' is the open-case default.", "Formula"),
    ("'Unknown' closure reason", "closure_reason", SUMIFS(("status", q("Close")), ("closure_reason", q("Unknown"))), f"{R['unknown']['top2_share']:.1%} of 'Unknown' closures are in two request types: Garbage Bin Request Case and Parking Enforcement Transfer Case. Almost all 'Unknown' closures are recorded on the same day the case opened.", "Kept as 'Unknown or N/A'. Reported by request type, not as a citywide rate only.", "Citywide unclear share is driven by two intake paths.", "Is 'Unknown' a default value for these intake paths?", "Formula"),
    ("Missing local area", "local_area", SUMIFS(("local_area", q("(Not recorded)"))), "Blank mainly for request types without a street location (feedback cases, bin requests, transfer cases).", "Labelled '(Not recorded)'. Not imputed.", "Area summaries describe located requests only.", "Which types are expected to carry a location?", "Formula"),
    ("Missing coordinates", "latitude, longitude", None, f"{Q['missing_coordinates']:,} records ({Q['missing_coordinates'] / N:.1%}). City metadata notes that some addresses are withheld and some do not geocode.", "Coordinates not used. Local-area names used instead.", "No point-level mapping (also a privacy choice).", "None for this study.", "Python"),
    ("Missing address", "address", None, f"{Q['missing_address']:,} records ({Q['missing_address'] / N:.1%}). Withheld for some request types by design.", "Presence flag only. Values removed from all outputs.", "None.", "None.", "Python"),
    ("Obsolete 'ZZ OLD' labels", "department, request type", SUMIFS(("service_request_type", q("ZZ*"))), "Three obsolete request types still received records in 2025 (for example 'ZZ OLD - Landscape Request Case').", "Kept with original labels.", "Negligible.", "Why do obsolete types still receive new cases?", "Formula"),
    ("Administrative/internal-sounding types", "service_request_type", SUMIFS(("admin_sounding", q("Yes"))), f"Rule: {ADMIN_RULE_TEXT} Four types flagged. 'Disposal Facility - Transfer Station Inquiry Case' is deliberately not flagged ('Transfer' names a facility).", "Sensitivity analysis: all records vs excluding flagged types.", "Flagged types hold 42% of 'Unknown' and 37% of auto-closed cases.", "Confirm whether these are resident-initiated.", "Formula"),
    ("No stable case identifier", "all", None, "No published field uniquely identifies a case.", "Cannot link referrals, reopenings or true duplicates.", "Limits process tracing and duplicate detection.", "Could an anonymised case ID be published?", "Python"),
    ("Exact duplicate rows", "all published fields", None, f"{Q['exact_duplicate_rows_all_columns']} exact duplicates across all published fields.", "None required.", "None.", "None.", "Python"),
    ("Duplicate-looking rows", "all fields except last modified", SUMIFS(("duplicate_looking", q("Yes"))), f"{Q['duplicate_looking_groups']:,} groups; mostly Garbage Bin and Green Bin Request Cases. {DB['timestamps_with_zero_seconds_share']:.0%} of opening timestamps end in :00 seconds, so coincidences are possible.", "Flagged only. Not removed.", "Demand counts may include multi-item or repeat requests.", "Does one resident request create one case per bin?", "Formula"),
    ("Coverage", "open date", None, f"All 12 months present; {R['profile']['distinct_open_dates']} distinct local opening dates. Monthly volume ranges {R['demand']['month_min']['count']:,}-{R['demand']['month_max']['count']:,}.", "None required.", "One year only: no seasonal claims.", "None.", "Python"),
    ("Source refresh and censoring", "all", None, f"The portal refreshes records from 17 Aug 2022 daily, so values can change. Latest close date in the extract: {Q['max_close_date']}. 2025 cases had 8-20 months to close before extraction.", "Analysis tied to the saved extract and its date.", "Right-censoring of durations is small but non-zero (open cases).", "None.", "Python"),
]
r = 5
for i, (check, fields, formula, finding, rule, effect, validate, by) in enumerate(dq, start=1):
    put(ws, f"A{r}", f"DQ{i:02d}", f_base, border=b_all, align=wrap)
    put(ws, f"B{r}", check, f_bold, border=b_all, align=wrap)
    put(ws, f"C{r}", fields, f_base, border=b_all, align=wrap)
    if formula:
        put(ws, f"D{r}", formula, f_base, INT, border=b_all)
        put(ws, f"E{r}", f"=D{r}/{N}" if i != 1 else f"=D{r}/D{r}", f_base, PCT, border=b_all)
    else:
        put(ws, f"D{r}", "", f_base, border=b_all)
        put(ws, f"E{r}", "", f_base, border=b_all)
    for colL, v in zip("FGHIJ", (finding, rule, effect, validate, by)):
        put(ws, f"{colL}{r}", v, f_base, border=b_all, align=wrap)
    r += 1
put(ws, f"A{r + 1}", f"Percentages use the cohort total of {N:,} records as the denominator.", f_note)
ws.freeze_panes = "C5"

# ------------------------------------------------------------------ Analysis
ws = S["Analysis"]
for L, w in zip("ABCDEFGHIJKLMN", [46, 14, 12, 14, 12, 14, 12, 14, 12, 14, 12, 14, 12, 14]):
    ws.column_dimensions[L].width = w
put(ws, "A1", "Analysis - every figure is a formula on the Power Query outputs", f_title)
put(ws, "A2", "Outcome shares use CLOSED cases as the denominator. Open cases are reported separately. Outcome groups are analyst-defined (see Closure_Outcome_Mapping).", f_note)
CHECKS = []   # (cell, expected, label)
CELLS = {}
r = 4

def section(title, note=None):
    global r
    put(ws, f"A{r}", title, f_h2, fill=fill_sec)
    for L in "BCDEFGHIJKLMN":
        ws[f"{L}{r}"].fill = fill_sec
    r += 1
    if note:
        put(ws, f"A{r}", note, f_note)
        r += 1

def check(label, formula):
    global r
    put(ws, f"A{r}", label, f_bold, fill=fill_check)
    put(ws, f"B{r}", formula, f_bold, fill=fill_check)
    CHECKS.append(f"Analysis!B{r}")
    r += 1

section("A. Headline counts")
put(ws, f"A{r}", "Total 2025 records"); put(ws, f"B{r}", SUMIFS(), fmt=INT); CELLS["total"] = f"B{r}"; r += 1
put(ws, f"A{r}", "Closed cases"); put(ws, f"B{r}", SUMIFS(("status", q("Close"))), fmt=INT); CELLS["closed"] = f"B{r}"; r += 1
put(ws, f"A{r}", "Open at extraction"); put(ws, f"B{r}", SUMIFS(("status", q("Open"))), fmt=INT); CELLS["open"] = f"B{r}"; r += 1
put(ws, f"A{r}", "Open share of all records"); put(ws, f"B{r}", f"={CELLS['open']}/{CELLS['total']}", fmt=PCT); CELLS["open_share"] = f"B{r}"; r += 1
put(ws, f"A{r}", "Expected total from the Python reference analysis", f_note); put(ws, f"B{r}", N, f_input, INT); exp_row = r; r += 1
check("Check: workbook total equals reference total", f"={CELLS['total']}=B{exp_row}")
check("Check: closed + open equals total", f"={CELLS['closed']}+{CELLS['open']}={CELLS['total']}")
r += 1

section("B. Closed-case outcome groups - all records vs excluding admin/internal-sounding types",
        "Sensitivity: the second pair of columns removes the four request types flagged by the name-based rule (see Data_Quality_Log DQ14).")
header_row(ws, r, 1, ["Outcome group (analyst-defined)", "Closed (all)", "Share (all)", "Closed (excl.)", "Share (excl.)", "Change (pts)"]); r += 1
b_start = r
for g in GROUPS:
    put(ws, f"A{r}", g)
    put(ws, f"B{r}", SUMIFS(("status", q("Close")), ("outcome_group", f"$A{r}")), fmt=INT)
    put(ws, f"C{r}", f"=B{r}/SUM($B${b_start}:$B${b_start + 5})", fmt=PCT)
    put(ws, f"D{r}", SUMIFS(("status", q("Close")), ("outcome_group", f"$A{r}"), ("admin_sounding", q("No"))), fmt=INT)
    put(ws, f"E{r}", f"=D{r}/SUM($D${b_start}:$D${b_start + 5})", fmt=PCT)
    put(ws, f"F{r}", f"=(E{r}-C{r})*100", fmt="+0.0;-0.0;0.0")
    CELLS[f"grp_all_{g[0]}"] = f"B{r}"; CELLS[f"grp_share_{g[0]}"] = f"C{r}"
    CELLS[f"grp_excl_{g[0]}"] = f"D{r}"; CELLS[f"grp_excl_share_{g[0]}"] = f"E{r}"
    r += 1
put(ws, f"A{r}", "Total closed", f_bold); put(ws, f"B{r}", f"=SUM(B{b_start}:B{r - 1})", f_bold, INT); put(ws, f"D{r}", f"=SUM(D{b_start}:D{r - 1})", f_bold, INT)
CELLS["closed_excl"] = f"D{r}"; b_tot = r; r += 1
ws.conditional_formatting.add(f"C{b_start}:C{b_start + 5}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2A78D6"))
check("Check: outcome groups sum to closed cases", f"=B{b_tot}={CELLS['closed']}")
r += 1

section("C. Closure-reason values (original, unmodified)")
header_row(ws, r, 1, ["Closure reason (original)", "Closed", "Share of closed", "Open", "Outcome group"]); r += 1
c_start = r
for reason in sorted(MAPPING, key=lambda k: -next(m["closed_2025"] for m in R["mapping"] if m["closure_reason_original"] == k)):
    put(ws, f"A{r}", reason, f_input)
    put(ws, f"B{r}", SUMIFS(("status", q("Close")), ("closure_reason", f"$A{r}")), fmt=INT)
    put(ws, f"C{r}", f"=B{r}/{CELLS['closed']}", fmt=PCT)
    put(ws, f"D{r}", SUMIFS(("status", q("Open")), ("closure_reason", f"$A{r}")), fmt=INT)
    put(ws, f"E{r}", f"=INDEX(tblClosureMapping[Outcome group (analyst-defined)],MATCH(A{r},tblClosureMapping[Closure reason (original)],0))")
    r += 1
check("Check: closure reasons sum to total records", f"=SUM(B{c_start}:B{r - 1})+SUM(D{c_start}:D{r - 1})={CELLS['total']}")
r += 1

section("D. Outcome mix by department - 12 largest departments by closed cases",
        "Department list chosen by the analyst (blue). Shares use each department's closed cases.")
header_row(ws, r, 1, ["Department", "Closed"] + [SHORT[g] for g in GROUPS]); r += 1
d_start = r
top12 = sorted(R["dept_variation"]["rows"], key=lambda x: -x["closed"])[:12]
for row in top12:
    put(ws, f"A{r}", row["department"], f_input)
    put(ws, f"B{r}", SUMIFS(("status", q("Close")), ("department", f"$A{r}")), fmt=INT)
    for j, g in enumerate(GROUPS):
        L = get_column_letter(3 + j)
        put(ws, f"{L}{r}", f"=IFERROR(SUMIFS({col('records')},{col('status')},\"Close\",{col('department')},$A{r},{col('outcome_group')},{q(g)})/$B{r},0)", fmt=PCT)
    r += 1
d_end = r - 1
CELLS["dept_range"] = (d_start, d_end)
ws.conditional_formatting.add(f"C{d_start}:C{d_end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2A78D6"))
ws.conditional_formatting.add(f"D{d_start}:D{d_end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="EB6834"))
put(ws, f"A{r}", "Share of all closed cases covered by these 12 departments", f_note)
put(ws, f"B{r}", f"=SUM(B{d_start}:B{d_end})/{CELLS['closed']}", fmt=PCT); r += 2

section("E. Recorded demand and outcomes by month (local opening month)")
header_row(ws, r, 1, ["Month", "Records", "Closed", "Service provided share", "Assigned / referred / continuing share", "Open"]); r += 1
e_start = r
months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
for mi, mn in enumerate(months, start=1):
    put(ws, f"A{r}", mn)
    put(ws, f"B{r}", SUMIFS(("open_month", str(mi))), fmt=INT)
    put(ws, f"C{r}", SUMIFS(("open_month", str(mi)), ("status", q("Close"))), fmt=INT)
    put(ws, f"D{r}", f"=IFERROR(SUMIFS({col('records')},{col('open_month')},{mi},{col('status')},\"Close\",{col('outcome_group')},{q(G1)})/C{r},0)", fmt=PCT)
    put(ws, f"E{r}", f"=IFERROR(SUMIFS({col('records')},{col('open_month')},{mi},{col('status')},\"Close\",{col('outcome_group')},{q(G2)})/C{r},0)", fmt=PCT)
    put(ws, f"F{r}", SUMIFS(("open_month", str(mi)), ("status", q("Open"))), fmt=INT)
    r += 1
e_end = r - 1
CELLS["month_range"] = (e_start, e_end)
check("Check: months sum to total records", f"=SUM(B{e_start}:B{e_end})={CELLS['total']}")
r += 1

section("F. Recorded demand by channel and outcome mix (channels with 1,000+ closed cases)")
header_row(ws, r, 1, ["Channel", "Records", "Share of records", "Closed", "Service provided", "Assigned / referred / continuing", "Unknown or N/A"]); r += 1
f_start = r
for ch in [x["value"] for x in R["demand"]["by_channel"]]:
    put(ws, f"A{r}", ch, f_input)
    put(ws, f"B{r}", SUMIFS(("channel", f"$A{r}")), fmt=INT)
    put(ws, f"C{r}", f"=B{r}/{CELLS['total']}", fmt=PCT)
    put(ws, f"D{r}", SUMIFS(("channel", f"$A{r}"), ("status", q("Close"))), fmt=INT)
    for L, g in zip("EFG", (G1, G2, G5)):
        put(ws, f"{L}{r}", f"=IF(D{r}<1000,\"n/a (<1,000)\",SUMIFS({col('records')},{col('channel')},$A{r},{col('status')},\"Close\",{col('outcome_group')},{q(g)})/D{r})", fmt=PCT)
    r += 1
check("Check: channels sum to total records", f"=SUM(B{f_start}:B{r - 1})={CELLS['total']}")
r += 1

section("G. Top 10 request types and demand concentration")
header_row(ws, r, 1, ["Request type", "Department", "Records", "Share of records", "Cumulative share"]); r += 1
g_start = r
for i, t in enumerate(R["demand"]["top_types"]):
    put(ws, f"A{r}", t["type"], f_input)
    put(ws, f"B{r}", t["department"], f_input)
    put(ws, f"C{r}", SUMIFS(("service_request_type", f"$A{r}"), ("department", f"$B{r}")), fmt=INT)
    put(ws, f"D{r}", f"=C{r}/{CELLS['total']}", fmt=PCT)
    put(ws, f"E{r}", f"=SUM($C${g_start}:C{r})/{CELLS['total']}", fmt=PCT)
    r += 1
CELLS["top5_share"] = f"E{g_start + 4}"; CELLS["top10_share"] = f"E{g_start + 9}"
put(ws, f"A{r}", "Share of records in the 5 largest departments")
put(ws, f"B{r}", "=" + "+".join(SUMIFS(("department", q(d["department"])))[1:] for d in R["demand"]["by_department_top10"][:5]), fmt=INT)
put(ws, f"C{r}", f"=B{r}/{CELLS['total']}", fmt=PCT); CELLS["top5_dept_share"] = f"C{r}"; r += 2

section("H. Where 'Unknown' closures come from")
header_row(ws, r, 1, ["Request type", "Closed as 'Unknown'", "Closed cases of this type", "Unknown rate within type", "Share of all 'Unknown'", "Admin/internal-sounding?"]); r += 1
h_start = r
for t in R["unknown"]["top_types"]:
    put(ws, f"A{r}", t["type"], f_input)
    put(ws, f"B{r}", SUMIFS(("service_request_type", f"$A{r}"), ("status", q("Close")), ("closure_reason", q("Unknown"))), fmt=INT)
    put(ws, f"C{r}", SUMIFS(("service_request_type", f"$A{r}"), ("status", q("Close"))), fmt=INT)
    put(ws, f"D{r}", f"=B{r}/C{r}", fmt=PCT)
    put(ws, f"E{r}", f"=B{r}/{SUMIFS(('status', q('Close')), ('closure_reason', q('Unknown')))[1:]}", fmt=PCT)
    put(ws, f"F{r}", f"=IF(SUMIFS({col('records')},{col('service_request_type')},$A{r},{col('admin_sounding')},\"Yes\")>0,\"Yes\",\"No\")")
    r += 1
put(ws, f"A{r}", "Top two request types' share of all 'Unknown' closures", f_bold)
put(ws, f"B{r}", f"=SUM(E{h_start}:E{h_start + 1})", f_bold, PCT); CELLS["unknown_top2"] = f"B{r}"; r += 2

section("I. Recorded demand by local area (aggregated; no addresses)")
header_row(ws, r, 1, ["Local area", "Records", "Share of records"]); r += 1
i_start = r
for a in R["demand"]["by_local_area"]:
    put(ws, f"A{r}", a["value"], f_input)
    put(ws, f"B{r}", SUMIFS(("local_area", f"$A{r}")), fmt=INT)
    put(ws, f"C{r}", f"=B{r}/{CELLS['total']}", fmt=PCT)
    r += 1
check("Check: local areas sum to total records", f"=SUM(B{i_start}:B{r - 1})={CELLS['total']}")
r += 1

section(f"J. Focus family: {FOCUS_FAMILY} ({fam_dept})",
        "Selected on evidence: volume, several meaningful outcome groups, low unclear share, comparable request types. See selection matrix below.")
header_row(ws, r, 1, ["Request type", "Records", "Closed", "Open"] + [SHORT[g] for g in GROUPS]); r += 1
j_start = r
for t in fam_types + ["All three types"]:
    put(ws, f"A{r}", t, f_input)
    crit_t = [("focus_family", q("Yes"))] + ([] if t == "All three types" else [("service_request_type", f"$A{r}")])
    put(ws, f"B{r}", SUMIFS(*crit_t), fmt=INT)
    put(ws, f"C{r}", SUMIFS(*(crit_t + [("status", q("Close"))])), fmt=INT)
    put(ws, f"D{r}", SUMIFS(*(crit_t + [("status", q("Open"))])), fmt=INT)
    for j, g in enumerate(GROUPS):
        L = get_column_letter(5 + j)
        put(ws, f"{L}{r}", "=IFERROR(" + SUMIFS(*(crit_t + [("status", q("Close")), ("outcome_group", q(g))]))[1:] + f"/C{r},0)", fmt=PCT)
    r += 1
j_end = r - 1
CELLS["fam_range"] = (j_start, j_end)
r += 1
put(ws, f"A{r}", "Focus-family selection matrix (from 03_analyze.py)", f_bold); r += 1
header_row(ws, r, 1, ["Candidate family", "Records", "Outcome groups >=5%", "Service provided", "Assigned / referred / continuing", "Unknown or N/A", "Criteria met (of 4)"]); r += 1
for s in R["focus_selection"]:
    put(ws, f"A{r}", s["family"] + (" (selected)" if s["family"] == FOCUS_FAMILY else ""), f_bold if s["family"] == FOCUS_FAMILY else f_base)
    put(ws, f"B{r}", s["records"], f_input, INT)
    put(ws, f"C{r}", s["groups_ge5pct"], f_input)
    put(ws, f"D{r}", s["service_provided"], f_input, PCT)
    put(ws, f"E{r}", s["continuing"], f_input, PCT)
    put(ws, f"F{r}", s["unclear"], f_input, PCT)
    put(ws, f"G{r}", s["criteria_met"], f_input)
    r += 1
put(ws, f"A{r}", "Criteria: 5,000+ records; 3+ outcome groups at 5% or more; unclear outcomes 5% or less; 2-5 comparable request types. Service-delivery connection judged qualitatively.", f_note); r += 2

section("K. Recorded calendar days to closure (nearest-rank percentiles, eligible closed cases)",
        "Computed from Source_Durations. Percentiles are shown only where n >= 30. Durations are calendar days between the local opening date and the recorded close date.")
header_row(ws, r, 1, ["Scope / request type / outcome", "Population", "n", "Median (p50)", "p75", "p90"]); r += 1
DUR_ROWS = []
def dur_row(scope, rtype, key, pop, label):
    global r
    put(ws, f"A{r}", label)
    put(ws, f"B{r}", pop)
    crit = f"{dcol('population')},$B{r},{dcol('scope')},{q(scope)},{dcol('request_type')},{q(rtype)},{dcol('outcome_key')},{q(key)}"
    put(ws, f"C{r}", f"=SUMIFS({dcol('records')},{crit})", fmt=INT)
    for L, p in zip("DEF", (0.5, 0.75, 0.9)):
        put(ws, f"{L}{r}", f"=IF(C{r}<30,\"n<30\",_xlfn.MINIFS({dcol('days_to_close')},{crit},{dcol('cum_records')},\">=\"&{p}*C{r}))")
    DUR_ROWS.append((scope, rtype, key, pop, r))
    r += 1
for pop in ("All published records", "Excluding admin/internal-sounding"):
    dur_row("Citywide", "All types", "All outcomes (mixed)", pop, "Citywide - all outcomes mixed")
    for g in GROUPS:
        dur_row("Citywide", "All types", g, pop, "Citywide - " + SHORT[g])
r += 1
for t in ["All three types"] + fam_types:
    for key, lab in [("All outcomes (mixed)", "all outcomes mixed"), (G1, "Service provided"), (G2, "Assigned / referred / continuing"),
                     ("Further action has been planned", "of which: Further action has been planned"),
                     ("Referred to another service group", "of which: Referred to another service group"), (G3, "No service or no action")]:
        dur_row("Focus family", t, key, "All published records", f"{t.replace(' Case', '')} - {lab}")
r += 1
put(ws, f"A{r}", "Duration eligibility (all records)", f_bold); r += 1
for lab, crit in (("Eligible closed cases", [("duration_exclusion_reason", '""')]),
                  ("Excluded: open at extraction", [("duration_exclusion_reason", q("Open at extraction (no close date)"))]),
                  ("Excluded: close date before local opening date", [("duration_exclusion_reason", q("Close date before local open date"))])):
    put(ws, f"A{r}", lab)
    put(ws, f"B{r}", SUMIFS(*crit), fmt=INT)
    CELLS["elig_" + lab.split(":")[0].split()[0].lower()] = f"B{r}"
    r += 1
check("Check: eligible + excluded equals total", f"=SUM(B{r - 3}:B{r - 1})={CELLS['total']}")
check("Check: eligible equals n in the citywide mixed duration row", f"=B{r - 4}=C{DUR_ROWS[0][4]}")
r += 1

section("L. Reconciliation summary")
put(ws, f"A{r}", "All checks TRUE", f_bold, fill=fill_check)
put(ws, f"B{r}", "=AND(" + ",".join(CHECKS + [map_check]) + ")", f_bold, fill=fill_check)
CELLS["all_checks"] = f"B{r}"
ws.freeze_panes = "B4"

# ------------------------------------------------------------------ Dashboard
ws = S["Dashboard"]
for L in "ABCDEFGHIJKLMNOPQRSTUVWX":
    ws.column_dimensions[L].width = 11
put(ws, "A1", "3-1-1 closure outcomes - management view (2025 cohort)", f_title)
put(ws, "A2", "Independent outside-in analysis of public data. Outcome groups are analyst-defined. All values are formulas linked to the Analysis sheet.", f_note)
tiles = [
    ("Requests opened in 2025", f"=Analysis!{CELLS['total']}", INT),
    ("Closed cases", f"=Analysis!{CELLS['closed']}", INT),
    ("Closed as 'Service provided'", f"=Analysis!{CELLS['grp_share_1']}", PCT),
    ("Closed: assigned / referred / continuing", f"=Analysis!{CELLS['grp_share_2']}", PCT),
    ("Closed as 'Unknown'", f"=Analysis!{CELLS['grp_share_5']}", PCT),
    ("Open at extraction", f"=Analysis!{CELLS['open']}", INT),
]
fill_tile = PatternFill("solid", fgColor="F3F6F9")
for i, (lab, fml, fmt) in enumerate(tiles):
    c0 = 1 + i * 4
    L0, L1 = get_column_letter(c0), get_column_letter(c0 + 2)
    ws.merge_cells(f"{L0}4:{L1}4"); ws.merge_cells(f"{L0}5:{L1}5")
    put(ws, f"{L0}4", lab, Font(name=FONT, size=9, color="52514E"), fill=fill_tile, align=Alignment(wrap_text=True))
    put(ws, f"{L0}5", fml, Font(name=FONT, size=18, bold=True, color="17324D"), fmt, fill=fill_tile)
    for L in (L0, get_column_letter(c0 + 1), L1):
        ws[f"{L}4"].fill = fill_tile; ws[f"{L}5"].fill = fill_tile
ws.row_dimensions[4].height = 26; ws.row_dimensions[5].height = 30
put(ws, "A6", "Denominator for percentages: closed cases. 'Service provided' is a recorded closure reason; it is not verified service completion.", f_note)
put(ws, "Q6", "Reconciliation checks:", f_bold); put(ws, "T6", f"=IF(Analysis!{CELLS['all_checks']},\"All TRUE\",\"CHECK FAILED\")", f_bold)

COLORS = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]
def finish(chart):
    # openpyxl 3.1 writes axes as deleted unless told otherwise -> no category labels in Excel.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    if chart.type == "bar":
        # categories run top-to-bottom (maxMin); put the value axis at the bottom
        chart.y_axis.crosses = "max"
    # reserve space for title and legend instead of drawing them over the plot
    if chart.title is not None:
        chart.title.overlay = False
    if chart.legend is not None:
        chart.legend.position = "b"
        chart.legend.overlay = False
    return chart
wsA = S["Analysis"]
# Chart 1: outcome groups (all vs excl)
ch = BarChart(); ch.type = "bar"; ch.style = 10
ch.title = "Closed cases by outcome group: all records vs excluding admin/internal-sounding"
ch.y_axis.title = "Share of closed cases"; ch.y_axis.numFmt = "0%"; ch.y_axis.majorGridlines = None
ch.add_data(Reference(wsA, min_col=3, min_row=b_start - 1, max_row=b_start + 5), titles_from_data=True)
ch.add_data(Reference(wsA, min_col=5, min_row=b_start - 1, max_row=b_start + 5), titles_from_data=True)
ch.set_categories(Reference(wsA, min_col=1, min_row=b_start, max_row=b_start + 5))
ch.series[0].graphicalProperties.solidFill = "2A78D6"; ch.series[1].graphicalProperties.solidFill = "B7D3F6"
ch.x_axis.scaling.orientation = "maxMin"
ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False, showLeaderLines=False); ch.dataLabels.numFmt = "0.0%"
ch.height, ch.width = 8.5, 17
ws.add_chart(finish(ch), "A8")
# Chart 2: department mix (percent stacked)
ch = BarChart(); ch.type = "bar"; ch.grouping = "percentStacked"; ch.overlap = 100; ch.gapWidth = 60
ch.title = "Outcome mix differs sharply by department (12 largest, closed cases)"
for j in range(6):
    ch.add_data(Reference(wsA, min_col=3 + j, min_row=d_start - 1, max_row=d_end), titles_from_data=True)
    ch.series[j].graphicalProperties.solidFill = COLORS[j]
    ch.series[j].graphicalProperties.line.solidFill = "FFFFFF"
ch.set_categories(Reference(wsA, min_col=1, min_row=d_start, max_row=d_end))
ch.y_axis.numFmt = "0%"; ch.x_axis.scaling.orientation = "maxMin"; ch.legend.position = "b"
ch.height, ch.width = 11, 24
ws.add_chart(finish(ch), "L8")
# Chart 3: monthly demand
ch = BarChart(); ch.type = "col"
ch.title = "Recorded requests by local opening month, 2025"
ch.add_data(Reference(wsA, min_col=2, min_row=e_start - 1, max_row=e_end), titles_from_data=True)
ch.set_categories(Reference(wsA, min_col=1, min_row=e_start, max_row=e_end))
ch.series[0].graphicalProperties.solidFill = "2A78D6"; ch.legend = None; ch.y_axis.numFmt = "#,##0"
ch.height, ch.width = 7.5, 17
ws.add_chart(finish(ch), "A26")
# Chart 4: focus family mix
ch = BarChart(); ch.type = "bar"; ch.grouping = "percentStacked"; ch.overlap = 100; ch.gapWidth = 60
ch.title = f"{FOCUS_FAMILY}: outcome mix by request type"
for j in range(6):
    ch.add_data(Reference(wsA, min_col=5 + j, min_row=j_start - 1, max_row=j_end), titles_from_data=True)
    ch.series[j].graphicalProperties.solidFill = COLORS[j]
    ch.series[j].graphicalProperties.line.solidFill = "FFFFFF"
ch.set_categories(Reference(wsA, min_col=1, min_row=j_start, max_row=j_end))
ch.y_axis.numFmt = "0%"; ch.x_axis.scaling.orientation = "maxMin"; ch.legend.position = "b"
ch.height, ch.width = 8, 24
ws.add_chart(finish(ch), "L31")

# ------------------------------------------------------------------ Assumptions_and_Limitations
ws = S["Assumptions_and_Limitations"]
ws.column_dimensions["A"].width = 6
for L, w in zip("BCDE", [40, 60, 50, 50]):
    ws.column_dimensions[L].width = w
put(ws, "A1", "Assumptions and limitations", f_title)
put(ws, "A2", "What the public data can and cannot support. Items marked 'Validate' need confirmation with City staff before any operational conclusion.", f_note)
header_row(ws, 4, 1, ["#", "Assumption or limitation", "Why it matters", "How this study handled it", "What to validate / with whom"])
al = [
    ("Outside-in only", "No interviews, internal documents or system access were used.", "Findings describe recorded data, not operations. Process maps are hypotheses.", "Validate: service owners, 3-1-1 Contact Centre, CRM/reporting team."),
    ("Closure reasons describe the case record, not necessarily the service", "'Closed' can mean handed off, planned or routed.", "Analyst-defined outcome groups; original values kept.", "Validate: the operational meaning of each of the 17 values."),
    ("'Service provided' taken at face value", "It may mean a response was given rather than work completed.", "Labelled 'service recorded as provided'.", "Validate: definition per department."),
    ("Close date is a local calendar date", "Affects same-day and one-day durations.", "Evidence-based assumption (DQ03); durations reported in calendar days.", "Validate: CRM close-date stamping."),
    ("Durations are recorded calendar days to closure", "Not elapsed, service-delivery or resolution time; handoff closures can be fast while work continues.", "Percentiles reported within outcome group and request type; no cross-service comparison.", "Validate: whether downstream completion dates exist."),
    ("No published service standards used", "No official response-time target was located in the public sources consulted (open-data portal and the City's Van311 page).", "Percentile-based descriptions only; no 'late' cases.", "Validate: any internal or published service standards."),
    ("Admin/internal-sounding flag is name-based", "Names may not reflect how cases are created.", "Sensitivity analysis instead of exclusion.", "Validate: which types are resident-initiated."),
    ("No case identifier", "Referrals, reopenings and duplicates cannot be traced.", "Duplicate-looking rows flagged, not removed.", "Validate: whether an anonymised ID could be published."),
    ("Last known status only", "Reopened or re-closed cases look like single events.", "Open cases reported separately; 33 status/reason conflicts flagged.", "Validate: reopening rules."),
    ("One calendar year", "Monthly variation cannot be called seasonal.", "No seasonal claims.", "Extend to 2023-2024 for trend work."),
    ("Source data changes", "Daily refresh from 17 Aug 2022 means later extracts can differ.", "All figures tied to the extraction timestamp on README.", "Re-run the scripts and refresh to update."),
    ("Location fields partly withheld", "Area analysis covers located requests only.", "Local-area summaries only; no point data published.", "None required for this study."),
    ("Outcome-group boundaries are judgement calls", "E.g., 'Alternate Service Required' could be redirection or a different City service.", "Placed in 'Other - requires review'.", "Validate: ambiguous values in a mapping workshop."),
]
r = 5
for i, (a, why, how, val) in enumerate(al, start=1):
    for L, v in zip("ABCDE", (i, a, why, how, val)):
        put(ws, f"{L}{r}", v, f_bold if L == "B" else f_base, align=wrap, border=b_all)
    r += 1

wb.save(STAGE)
# Save cell map for the reconciliation step
(ROOT / "build" / "workbook_cells.json").write_text(json.dumps({"cells": CELLS, "checks": CHECKS, "map_check": map_check,
    "dur_rows": DUR_ROWS, "b_start": b_start, "d_range": [d_start, d_end], "j_range": [j_start, j_end],
    "e_range": [e_start, e_end], "h_start": h_start}, indent=1))
print("stage-1 workbook saved:", STAGE)
