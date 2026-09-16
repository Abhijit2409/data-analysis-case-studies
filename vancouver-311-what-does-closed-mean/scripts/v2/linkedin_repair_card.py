"""Build the third LinkedIn image (street and sidewalk repair outcome mix + durations) from the workbook's saved values and render it with Edge."""
import pathlib
import subprocess
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
WB = ROOT / "outputs" / "Vancouver311_Closure_Outcomes_Workbook.xlsx"
OUT_DIR = ROOT / "outputs" / "v2_redesign"
HTML = OUT_DIR / "LinkedIn_Repair_Comparison_source.html"
PNG = OUT_DIR / "LinkedIn_Repair_Comparison.png"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

NAVY, TEAL, SLATE, LIGHT = "#1D3557", "#2A9D8F", "#8D99AE", "#D5DAE0"

ws = load_workbook(WB, data_only=True, read_only=True)["Analysis"]
cell = lambda ref: ws[ref].value
assert cell("A142").startswith("J. Focus family: Street and sidewalk repair")
assert [cell(f"{c}144") for c in "CEFGHIJ"] == ["Closed", "Service provided", "Assigned / referred / continuing", "No service or no action",
                                               "Insufficient info / unreachable", "Unknown or N/A", "Other - review"]
LABELS = {"Pothole Case": "Pothole", "Street Repair Case": "Street repair", "Sidewalk Repair Case": "Sidewalk repair", "All three types": "All three types"}

rows = []
for r in range(145, 149):
    name = cell(f"A{r}")
    sp, ho, ns = cell(f"E{r}"), cell(f"F{r}"), cell(f"G{r}")
    rest = cell(f"H{r}") + cell(f"I{r}") + cell(f"J{r}")
    assert abs(sp + ho + ns + rest - 1) < 1e-9
    rows.append({"label": LABELS[name], "closed": cell(f"C{r}"), "open": cell(f"D{r}"), "sp": sp, "ho": ho, "ns": ns, "rest": rest})
pothole, street, sidewalk, total = rows
assert f"{pothole['sp']:.0%}" == "77%" and f"{street['sp']:.0%}" == "38%" and total["open"] == 140

assert cell("A198") == "Sidewalk Repair - Service provided" and cell("A200") == "Sidewalk Repair - of which: Further action has been planned"
assert cell("B198") == cell("B200") == "All published records"
dur_sp = {"n": cell("C198"), "p50": cell("D198")}
dur_fa = {"n": cell("C200"), "p50": cell("D200")}
assert (dur_sp["n"], dur_sp["p50"], dur_fa["n"], dur_fa["p50"]) == (904, 5, 551, 9)
assert cell("A152") == "Street and sidewalk repair (selected)" and cell("G152") == 4
sel_ho = cell("E152")

BAR_W = 928
def seg(share, colour, text_colour):
    w = share * BAR_W
    txt = f"{share * 100:.0f}%" if share >= 0.08 else ""
    return f'<div class="seg" style="width:{w:.1f}px;background:{colour};color:{text_colour}">{txt}</div>'

def row(d, cls=""):
    return (f'<div class="row {cls}"><div class="top"><span class="name">{d["label"]}</span><span class="cnt">{d["closed"]:,} closed</span></div>'
            f'<div class="bar">{seg(d["sp"], NAVY, "#FFFFFF")}{seg(d["ho"], TEAL, "#FFFFFF")}{seg(d["ns"], SLATE, NAVY)}{seg(d["rest"], LIGHT, NAVY)}</div></div>')

page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Street and sidewalk repair card</title>
<style>
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; background: #FFFFFF; }}
body {{ width: 1080px; height: 1350px; overflow: hidden; font-family: "Segoe UI", Arial, sans-serif; color: #1F2933; }}
.card {{ position: relative; height: 1350px; padding: 58px 76px 0; }}
.kicker {{ font-size: 20px; letter-spacing: 2.4px; text-transform: uppercase; color: #1F7A70; font-weight: 600; margin: 0 0 12px; }}
h1 {{ font-size: 52px; line-height: 1.12; color: {NAVY}; margin: 0 0 16px; font-weight: 700; letter-spacing: -0.5px; }}
.dek {{ font-size: 24px; line-height: 1.4; color: #3E4C59; margin: 0 0 24px; }}
.legend {{ display: flex; flex-wrap: wrap; gap: 6px 22px; font-size: 19px; color: #3E4C59; margin: 0 0 20px; }}
.legend i {{ display: inline-block; width: 18px; height: 18px; border-radius: 3px; margin-right: 7px; vertical-align: -3px; }}
.row {{ margin: 0 0 16px; }}
.row.total {{ border-top: 1.5px solid #E4E7EB; padding-top: 14px; }}
.top {{ display: flex; justify-content: space-between; align-items: baseline; margin: 0 0 5px; }}
.name {{ font-size: 25px; font-weight: 700; color: {NAVY}; }}
.cnt {{ font-size: 20px; color: #616E7C; }}
.bar {{ display: flex; height: 44px; border-radius: 3px; overflow: hidden; }}
.seg {{ height: 44px; font-size: 21px; font-weight: 700; line-height: 44px; padding-left: 10px; white-space: nowrap; overflow: hidden; }}
.note {{ font-size: 18px; color: #616E7C; line-height: 1.4; margin: 2px 0 0; }}
h2 {{ font-size: 27px; color: {NAVY}; margin: 30px 0 12px; font-weight: 700; }}
.stats {{ display: grid; grid-template-columns: 1fr 1fr; gap: 22px; }}
.stat {{ border-top: 5px solid {NAVY}; background: #F5F7FA; padding: 14px 20px 14px; }}
.stat.teal {{ border-top-color: {TEAL}; }}
.stat .v {{ font-size: 50px; font-weight: 700; color: {NAVY}; line-height: 1; }}
.stat.teal .v {{ color: #1F7A70; }}
.stat .v small {{ font-size: 22px; font-weight: 600; color: #3E4C59; margin-left: 6px; }}
.stat .l {{ font-size: 21px; color: #1F2933; margin-top: 6px; line-height: 1.3; }}
.stat .d {{ font-size: 18px; color: #616E7C; margin-top: 3px; }}
footer {{ position: absolute; left: 76px; right: 76px; bottom: 30px; display: flex; justify-content: space-between; gap: 30px;
  font-size: 17px; color: #7B8794; border-top: 1.5px solid #E4E7EB; padding-top: 12px; }}
</style></head>
<body><div class="card">
<div class="kicker">Vancouver 3-1-1 &middot; Street and sidewalk repair &middot; 2025</div>
<h1>Potholes: {pothole['sp']:.0%} &ldquo;Service provided.&rdquo; Street repair: {street['sp']:.0%}.</h1>
<p class="dek">Same service family, same &ldquo;closed&rdquo; status. What the closure reason records depends on the request type.</p>
<div class="legend">
 <span><i style="background:{NAVY}"></i>Service provided</span>
 <span><i style="background:{TEAL}"></i>Handed off or planned</span>
 <span><i style="background:{SLATE}"></i>No service or no action</span>
 <span><i style="background:{LIGHT}"></i>Unknown or other</span></div>
{row(pothole)}
{row(street)}
{row(sidewalk)}
{row(total, "total")}
<p class="note">Share of each request type&rsquo;s closed cases; {total['open']} open cases excluded.</p>

<h2>Sidewalk repair: planned work closes later</h2>
<div class="stats">
 <div class="stat"><div class="v">{dur_sp['p50']}<small>days (median)</small></div><div class="l">&ldquo;Service provided&rdquo;</div><div class="d">n = {dur_sp['n']:,} closed cases</div></div>
 <div class="stat teal"><div class="v">{dur_fa['p50']}<small>days (median)</small></div><div class="l">&ldquo;Further action has been planned&rdquo;</div><div class="d">n = {dur_fa['n']:,} closed cases</div></div>
</div>
<p class="note" style="margin-top:10px">Recorded calendar days from opening to closure, not time to complete the work. Selected for the pilot: it met all four
of the analyst&rsquo;s selection tests, with {sel_ho:.0%} of closures handed off or planned.</p>

<footer><span>Independent analysis of City of Vancouver open data. Not affiliated with or endorsed by the City.</span><span>Abhijit Mishra</span></footer>
</div></body></html>
"""
HTML.write_text(page, encoding="utf-8")
subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=2",
                "--window-size=1080,1350", f"--screenshot={PNG}", HTML.resolve().as_uri()], check=True, timeout=120)
print("rows:", [(d["label"], d["closed"], f'{d["sp"]:.1%}', f'{d["ho"]:.1%}', f'{d["ns"]:.1%}', f'{d["rest"]:.1%}') for d in rows])
print("durations:", dur_sp, dur_fa, "| selection handoff share:", f"{sel_ho:.1%}")
print("wrote", PNG, PNG.stat().st_size)
