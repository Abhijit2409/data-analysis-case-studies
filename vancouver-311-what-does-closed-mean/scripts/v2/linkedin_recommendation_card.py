"""Build the fourth LinkedIn image (recommendation, proposed flow, applied vs needs validation) and render it with Edge.

All wording is taken from pages 3-4 of the v2 report; the only figure (17 closure reasons) is checked against the workbook mapping sheet.
"""
import pathlib
import subprocess
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
WB = ROOT / "outputs" / "Vancouver311_Closure_Outcomes_Workbook.xlsx"
OUT_DIR = ROOT / "outputs" / "v2_redesign"
HTML = OUT_DIR / "LinkedIn_Recommendation_source.html"
PNG = OUT_DIR / "LinkedIn_Recommendation.png"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
NAVY, TEAL = "#1D3557", "#2A9D8F"

ws = load_workbook(WB, data_only=True, read_only=True)["Closure_Outcome_Mapping"]
reasons = [ws[f"A{r}"].value for r in range(5, 22)]
assert len(reasons) == 17 and all(reasons) and ws["A22"].value in (None, "")

FLOW = [("Preserve", "Keep the original status and closure reason"),
        ("Group", "Apply agreed outcome definitions"),
        ("Separate", "Show open and unclear cases alongside"),
        ("Review", "Service owner checks exceptions monthly"),
        ("Publish", "With definitions and the data date")]
APPLIED = ["Outcome shares use closed cases; open cases shown separately.",
           f"All {len(reasons)} original closure reasons kept and mapped.",
           "Durations compared within request type and outcome."]
PROPOSED = ["Agreed definitions for each closure reason.",
            "An outcome-split view for the pilot service.",
            "A way to link completion of planned or referred work to the case."]

steps = "\n".join(f'<div class="step"><div class="num">{i}</div><div class="st">{t}</div><div class="sd">{d}</div></div>'
                  for i, (t, d) in enumerate(FLOW, 1))
ul = lambda items: "<ul>" + "".join(f"<li>{x}</li>" for x in items) + "</ul>"

page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Recommendation card</title>
<style>
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; background: #FFFFFF; }}
body {{ width: 1080px; height: 1350px; overflow: hidden; font-family: "Segoe UI", Arial, sans-serif; color: #1F2933; }}
.card {{ position: relative; height: 1350px; padding: 58px 76px 0; }}
.kicker {{ font-size: 20px; letter-spacing: 2.4px; text-transform: uppercase; color: #1F7A70; font-weight: 600; margin: 0 0 12px; }}
h1 {{ font-size: 54px; line-height: 1.12; color: {NAVY}; margin: 0 0 22px; font-weight: 700; letter-spacing: -0.5px; }}
.rec {{ background: {NAVY}; color: #FFFFFF; border-radius: 4px; padding: 20px 26px; margin: 0 0 34px; }}
.rec .k {{ font-size: 17px; letter-spacing: 2px; text-transform: uppercase; color: #B8DDD8; font-weight: 600; }}
.rec .h {{ font-size: 29px; font-weight: 700; line-height: 1.25; margin: 6px 0 6px; }}
.rec p {{ font-size: 21px; line-height: 1.4; margin: 0; color: #E4E7EB; }}
h2 {{ font-size: 27px; color: {NAVY}; margin: 0 0 14px; font-weight: 700; }}
.flow {{ position: relative; margin: 0 0 34px; }}
.flow::before {{ content: ""; position: absolute; left: 23px; top: 24px; bottom: 24px; width: 3px; background: #B8DDD8; }}
.step {{ position: relative; display: grid; grid-template-columns: 48px 150px 1fr; align-items: center; gap: 0 20px; margin: 0 0 12px; }}
.num {{ width: 48px; height: 48px; border-radius: 50%; background: {TEAL}; color: #FFFFFF; font-size: 23px; font-weight: 700;
  display: flex; align-items: center; justify-content: center; }}
.st {{ font-size: 25px; font-weight: 700; color: {NAVY}; }}
.sd {{ font-size: 23px; color: #3E4C59; }}
.two {{ display: grid; grid-template-columns: 1fr 1fr; gap: 22px; }}
.box {{ padding: 16px 22px 8px; border-top: 5px solid {TEAL}; background: #EEF6F5; }}
.box.grey {{ border-top-color: #8D99AE; background: #F5F7FA; }}
.box h3 {{ font-size: 23px; color: {NAVY}; margin: 0 0 8px; line-height: 1.25; }}
.box ul {{ margin: 0; padding-left: 24px; }}
.box li {{ font-size: 20px; line-height: 1.38; margin: 0 0 9px; color: #1F2933; }}
.note {{ font-size: 19px; color: #616E7C; line-height: 1.4; margin: 18px 0 0; }}
footer {{ position: absolute; left: 76px; right: 76px; bottom: 30px; display: flex; justify-content: space-between; gap: 30px;
  font-size: 17px; color: #7B8794; border-top: 1.5px solid #E4E7EB; padding-top: 12px; }}
</style></head>
<body><div class="card">
<div class="kicker">Vancouver 3-1-1 &middot; Recommendation &middot; Independent study</div>
<h1>Make closure outcomes visible in reporting</h1>
<div class="rec"><div class="k">Recommended pilot</div>
 <div class="h">Outcome-aware closure reporting for street and sidewalk repair</div>
 <p>Uses fields already recorded and needs no system change. Outcomes are reported alongside the closure count, not instead of it.</p></div>

<h2>Proposed reporting flow</h2>
<div class="flow">
{steps}
</div>

<div class="two">
 <div class="box"><h3>Already applied in this study</h3>{ul(APPLIED)}</div>
 <div class="box grey"><h3>Proposed: needs validation with City staff</h3>{ul(PROPOSED)}</div>
</div>
<p class="note">Pilot: three monthly cycles, then expand, adjust or stop against agreed criteria. Targets are set only after baselines are measured.</p>

<footer><span>Independent analysis of City of Vancouver open data. Not affiliated with or endorsed by the City.</span><span>Abhijit Mishra</span></footer>
</div></body></html>
"""
HTML.write_text(page, encoding="utf-8")
subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=2",
                "--window-size=1080,1350", f"--screenshot={PNG}", HTML.resolve().as_uri()], check=True, timeout=120)
print("closure reasons in mapping:", len(reasons))
print("wrote", PNG, PNG.stat().st_size)
