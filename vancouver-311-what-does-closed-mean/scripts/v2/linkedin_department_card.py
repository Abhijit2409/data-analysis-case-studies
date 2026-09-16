"""Build the second LinkedIn image (department outcome mix) from the workbook's saved values and render it with Edge."""
import pathlib
import subprocess
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
WB = ROOT / "outputs" / "Vancouver311_Closure_Outcomes_Workbook.xlsx"
OUT_DIR = ROOT / "outputs" / "v2_redesign"
HTML = OUT_DIR / "LinkedIn_Department_Mix_source.html"
PNG = OUT_DIR / "LinkedIn_Department_Mix.png"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

NAVY, TEAL, GREY = "#1D3557", "#2A9D8F", "#D5DAE0"
SHORT_DEPT = {"Traffic and Electrical Operations and Design": "Traffic & Electrical Ops & Design",
              "Parking Enforcement and Operations": "Parking Enforcement & Operations",
              "Business and Election Services": "Business & Election Services"}

ws = load_workbook(WB, data_only=True, read_only=True)["Analysis"]
cell = lambda ref: ws[ref].value
assert cell("A48") == "Department" and cell("C48") == "Service provided" and cell("D48") == "Assigned / referred / continuing"
assert cell("A61") == "Share of all closed cases covered by these 12 departments"
assert cell("B5") == 269798 + 2782 and cell("B6") == 269798

rows = []
for r in range(49, 61):
    name = cell(f"A{r}")
    unit, dept = name.split(" - ", 1) if " - " in name else ("", name)
    dept = SHORT_DEPT.get(dept, dept)
    rows.append({"title": f"{dept} ({unit})" if unit else dept, "closed": cell(f"B{r}"),
                 "sp": cell(f"C{r}"), "ho": cell(f"D{r}")})
rows.sort(key=lambda d: d["sp"], reverse=True)
lo, hi, cover = rows[-1]["sp"], rows[0]["sp"], cell("B61")
assert f"{lo:.1%}" == "1.8%" and f"{hi:.1%}" == "86.7%" and f"{cover:.1%}" == "79.6%"

BAR_W = 928
def seg(share, colour, label):
    w = share * BAR_W
    txt = f"{share * 100:.0f}%" if label and share >= 0.08 else ""
    return f'<div class="seg" style="width:{w:.1f}px;background:{colour}">{txt}</div>'

row_html = "\n".join(
    f'<div class="row"><div class="top"><span class="name">{d["title"]}</span><span class="cnt">{d["closed"]:,} closed</span></div>'
    f'<div class="bar">{seg(d["sp"], NAVY, True)}{seg(d["ho"], TEAL, True)}{seg(1 - d["sp"] - d["ho"], GREY, False)}</div></div>'
    for d in rows)

page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Department outcome mix card</title>
<style>
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; background: #FFFFFF; }}
body {{ width: 1080px; height: 1350px; overflow: hidden; font-family: "Segoe UI", Arial, sans-serif; color: #1F2933; }}
.card {{ position: relative; height: 1350px; padding: 58px 76px 0; }}
.kicker {{ font-size: 20px; letter-spacing: 2.4px; text-transform: uppercase; color: #1F7A70; font-weight: 600; margin: 0 0 12px; }}
h1 {{ font-size: 50px; line-height: 1.12; color: {NAVY}; margin: 0 0 16px; font-weight: 700; letter-spacing: -0.5px; }}
.legend {{ display: flex; flex-wrap: nowrap; gap: 24px; font-size: 20px; color: #3E4C59; margin: 0 0 18px; align-items: center; white-space: nowrap; }}
.legend i {{ display: inline-block; width: 19px; height: 19px; border-radius: 3px; margin-right: 8px; vertical-align: -3px; }}
.legend .lead {{ color: #616E7C; }}
.row {{ margin: 0 0 9px; }}
.top {{ display: flex; justify-content: space-between; align-items: baseline; margin: 0 0 3px; }}
.name {{ font-size: 22px; font-weight: 700; color: {NAVY}; }}
.cnt {{ font-size: 19px; color: #616E7C; }}
.bar {{ display: flex; height: 30px; border-radius: 3px; overflow: hidden; }}
.seg {{ height: 30px; color: #FFFFFF; font-size: 19px; font-weight: 700; line-height: 30px; padding-left: 9px; white-space: nowrap; overflow: hidden; }}
.note {{ font-size: 19px; color: #616E7C; line-height: 1.4; margin: 14px 0 0; border-top: 1.5px solid #E4E7EB; padding-top: 12px; }}
footer {{ position: absolute; left: 76px; right: 76px; bottom: 30px; display: flex; justify-content: space-between; gap: 30px;
  font-size: 17px; color: #7B8794; border-top: 1.5px solid #E4E7EB; padding-top: 12px; }}
</style></head>
<body><div class="card">
<div class="kicker">Vancouver 3-1-1 &middot; Closed 2025 requests &middot; 12 largest departments</div>
<h1>&ldquo;Service provided&rdquo; ranges from {lo:.1%} to {hi:.1%} across departments</h1>
<div class="legend"><span class="lead">Share of closed cases:</span>
 <span><i style="background:{NAVY}"></i>Service provided</span>
 <span><i style="background:{TEAL}"></i>Handed off or planned</span>
 <span><i style="background:{GREY}"></i>All other outcomes</span></div>
{row_html}
<p class="note">These 12 departments hold {cover:.1%} of 269,798 closed cases. Not a performance ranking: departments handle different
request types, so one closure rate can&rsquo;t compare them without the outcome split.</p>
<footer><span>Independent analysis of City of Vancouver open data. Not affiliated with or endorsed by the City.</span><span>Abhijit Mishra</span></footer>
</div></body></html>
"""
HTML.write_text(page, encoding="utf-8")
subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=2",
                "--window-size=1080,1350", f"--screenshot={PNG}", HTML.resolve().as_uri()], check=True, timeout=120)
print("rows:", [(d["title"], d["closed"], f'{d["sp"]:.1%}', f'{d["ho"]:.1%}') for d in rows])
print("wrote", PNG, PNG.stat().st_size)
